"""Bulk student import: Excel template generation + parsing.

Only five fields are mandatory (admission_no, name, class, division,
parent_whatsapp) so office staff can fill a sheet quickly; everything else
is optional and defaults sensibly, matching the single-student add form.
"""
import io
import re

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

MANDATORY_FIELDS = ["admission_no", "name", "class", "division", "parent_whatsapp"]

FIELD_LABELS = {
    "admission_no": "admission_no*",
    "name": "name*",
    "class": "class*",
    "division": "division*",
    "parent_whatsapp": "parent_whatsapp*",
    "parent_name": "parent_name",
    "address": "address",
    "place": "place",
    "school_name": "school_name",
    "dob": "dob (YYYY-MM-DD)",
    "admission_date": "admission_date (YYYY-MM-DD)",
    "discount_amount": "discount_amount",
    "discount_reason": "discount_reason",
}

ALL_FIELDS = list(FIELD_LABELS.keys())


def build_template(class_names):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Students"

    headers = [FIELD_LABELS[f] for f in ALL_FIELDS]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    ws.append(
        [
            "BW101", "Amal Krishna", class_names[0] if class_names else "9", "A",
            "9847012345", "Suresh Kumar", "Omassery", "Omassery",
            "Govt. Higher Secondary School, Omassery", "2011-05-14", "2026-06-01", 0, "", "",
        ]
    )
    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(header) + 2)

    info = wb.create_sheet("Instructions")
    info.append(["Field", "Required?", "Notes"])
    for cell in info[1]:
        cell.font = Font(bold=True)
    info.append(["admission_no", "Yes", "Must be unique for every student"])
    info.append(["name", "Yes", "Student's full name"])
    info.append(["class", "Yes", f"One of: {', '.join(class_names)}"])
    info.append(["division", "Yes", "e.g. A, B - created automatically if it doesn't exist yet"])
    info.append(["parent_whatsapp", "Yes", "10-digit mobile number (country code optional)"])
    info.append(["parent_name", "No", "Parent / guardian name"])
    info.append(["address", "No", ""])
    info.append(["place", "No", "Village/town the student is from"])
    info.append(["school_name", "No", "The regular school the student studies in"])
    info.append(["dob", "No", "Format YYYY-MM-DD"])
    info.append(["admission_date", "No", "Format YYYY-MM-DD, defaults to today"])
    info.append(["discount_amount", "No", "Rupees, defaults to 0"])
    info.append(["discount_reason", "No", "e.g. Sibling discount"])
    info.column_dimensions["A"].width = 20
    info.column_dimensions["B"].width = 12
    info.column_dimensions["C"].width = 55

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _normalize_header(value):
    if value is None:
        return ""
    text = str(value).strip().lower()
    text = text.split("(")[0]  # drop "(YYYY-MM-DD)" hints
    text = text.replace("*", "").strip()
    text = text.replace(" ", "_")
    return text


def parse_upload(file_stream):
    """Parses an uploaded workbook into a list of row dicts with normalized
    keys (matching FIELD_LABELS keys) plus a "_row" spreadsheet row number.
    Blank rows are skipped. Raises ValueError on unreadable files."""
    try:
        wb = openpyxl.load_workbook(file_stream, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surface as a friendly upload error
        raise ValueError(f"Not a valid Excel (.xlsx) file: {exc}") from exc

    ws = wb["Students"] if "Students" in wb.sheetnames else wb.worksheets[0]

    try:
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    except StopIteration:
        return []
    headers = [_normalize_header(h) for h in header_row]

    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue
        record = {"_row": row_idx}
        for header, value in zip(headers, row):
            if header:
                record[header] = value
        rows.append(record)
    return rows


# ----------------------------------------------------- exam marks bulk entry
MARKS_SHEET = "Marks"


def build_marks_template(exam, students, existing):
    """A marks grid for one exam: a row per student, a column per subject.

    Pre-filled with whatever is already entered, so the same sheet works for a
    first entry and for a correction - and downloading it never loses marks
    somebody has already typed in.

    `existing` maps (student_id, exam_subject_id) -> marks.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = MARKS_SHEET

    exam_subjects = list(exam.exam_subjects)
    headers = ["admission_no", "name", "class"] + [
        f"{es.subject.name} (max {es.max_marks:g})" for es in exam_subjects
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for student in students:
        row = [
            student.admission_no,
            student.name,
            f"{student.school_class.name}-{student.division.name}",
        ]
        for es in exam_subjects:
            row.append(existing.get((student.id, es.id)))
        ws.append(row)

    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(header) + 2)
    ws.freeze_panes = "D2"

    info = wb.create_sheet("Instructions")
    info.append(["How to fill this sheet"])
    info[1][0].font = Font(bold=True)
    for line in [
        "",
        f"Exam: {exam.name} - Class {exam.school_class.name}",
        "",
        "Type each student's marks under the subject column.",
        "Leave a cell blank to leave that mark unchanged.",
        "Enter 0 for a student who sat the exam and scored nothing.",
        "Do not change the admission_no column - it identifies the student.",
        "Marks above the subject maximum are rejected and reported back to you.",
        "Extra rows for students not in this class are ignored.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 70

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def parse_marks_upload(file_stream):
    """Reads a filled marks grid.

    Returns (rows, subject_columns) where each row is
    {"_row": n, "admission_no": str, "marks": {subject_name: value}}.
    """
    try:
        wb = openpyxl.load_workbook(file_stream, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surface as a friendly upload error
        raise ValueError(f"Not a valid Excel (.xlsx) file: {exc}") from exc

    ws = wb[MARKS_SHEET] if MARKS_SHEET in wb.sheetnames else wb.worksheets[0]

    try:
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    except StopIteration:
        return [], []

    subject_columns = {}
    admission_col = None
    for index, raw in enumerate(header_row):
        if raw is None:
            continue
        header = str(raw).strip()
        if _normalize_header(header) == "admission_no":
            admission_col = index
        elif header.lower() not in ("name", "class"):
            # "Mathematics (max 100)" -> "Mathematics"
            subject_columns[index] = header.split("(")[0].strip()

    if admission_col is None:
        raise ValueError("The sheet has no admission_no column - use the downloaded template.")

    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue
        admission_no = row[admission_col] if admission_col < len(row) else None
        if admission_no in (None, ""):
            continue
        marks = {
            name: row[index]
            for index, name in subject_columns.items()
            if index < len(row) and row[index] not in (None, "")
        }
        rows.append(
            {"_row": row_idx, "admission_no": str(admission_no).strip(), "marks": marks}
        )
    return rows, list(subject_columns.values())


# ------------------------------------------------ bulk fee payment import
PAYMENTS_SHEET = "Payments"
INSTALLMENT_PATTERN = re.compile(r"^installment_(\d+)$")
DATE_PATTERN = re.compile(r"^date_(\d+)$")
MODE_PATTERN = re.compile(r"^mode_(\d+)$")


def build_payments_template(students, installments=6):
    """One row per student, with several installment columns.

    A family usually pays in parts across the year, so a sheet of one payment
    per row would mean hunting for the same student several times. Here each
    student appears once and their payments run across the row.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = PAYMENTS_SHEET

    headers = ["admission_no*", "name", "class", "course_fee", "discount", "paid_so_far", "balance_due"]
    for n in range(1, installments + 1):
        headers += [f"installment_{n}", f"date_{n}", f"mode_{n}"]

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for student in students:
        row = [
            student.admission_no,
            student.name,
            f"{student.school_class.name}-{student.division.name}",
            student.class_fee,
            student.discount_amount or 0,
            student.total_paid,
            student.pending_fee,
        ]
        row += [None] * (installments * 3)
        ws.append(row)

    for i, header in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(13, len(header) + 2)
    ws.freeze_panes = "H2"

    info = wb.create_sheet("Instructions")
    info.append(["How to fill this sheet"])
    info[1][0].font = Font(bold=True)
    for line in [
        "",
        "Each student has one row. Enter each payment they made in its own installment column.",
        "",
        "installment_1 / date_1 / mode_1  - the first payment and when it was made",
        "installment_2 / date_2 / mode_2  - the second, and so on across the row",
        "",
        "Leave a pair blank if there was no such payment.",
        "Dates are YYYY-MM-DD; blank means today.",
        "mode is Cash, UPI, Bank Transfer or Cheque - blank means Cash.",
        "",
        "course_fee, discount, paid_so_far and balance_due are shown for reference only;",
        "they are not read back, so changing them does nothing.",
        "",
        "The installments on a row must not add up to more than that student's balance.",
        "A payment identical to one already recorded (same student, amount and date) is",
        "skipped, so uploading the same sheet twice does not double the collection.",
        "Need more columns? Add installment_7 / date_7 / mode_7 and so on - they are read.",
    ]:
        info.append([line])
    info.column_dimensions["A"].width = 84

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def parse_payments_upload(file_stream):
    """Rows of {admission_no, payments: [{n, amount, date, mode}]}.

    Installment columns are discovered by name, so a sheet with extra columns
    added by hand is read the same as the downloaded one.
    """
    try:
        wb = openpyxl.load_workbook(file_stream, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surface as a friendly upload error
        raise ValueError(f"Not a valid Excel (.xlsx) file: {exc}") from exc

    ws = wb[PAYMENTS_SHEET] if PAYMENTS_SHEET in wb.sheetnames else wb.worksheets[0]

    try:
        header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    except StopIteration:
        return []
    headers = [_normalize_header(h) for h in header_row]

    if "admission_no" not in headers:
        raise ValueError("The sheet has no admission_no column - use the downloaded template.")
    admission_col = headers.index("admission_no")

    amounts, dates, modes = {}, {}, {}
    for index, header in enumerate(headers):
        for pattern, target in (
            (INSTALLMENT_PATTERN, amounts),
            (DATE_PATTERN, dates),
            (MODE_PATTERN, modes),
        ):
            match = pattern.match(header or "")
            if match:
                target[int(match.group(1))] = index

    if not amounts:
        raise ValueError(
            "The sheet has no installment_1 column - use the downloaded template."
        )

    rows = []
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row is None or all(cell in (None, "") for cell in row):
            continue

        def cell(index):
            return row[index] if index is not None and index < len(row) else None

        admission_no = cell(admission_col)
        if admission_no in (None, ""):
            continue

        payments = []
        for n in sorted(amounts):
            amount = cell(amounts.get(n))
            if amount in (None, ""):
                continue
            payments.append(
                {
                    "n": n,
                    "amount": amount,
                    "date": cell(dates.get(n)),
                    "mode": cell(modes.get(n)),
                }
            )

        if payments:
            rows.append(
                {"_row": row_idx, "admission_no": str(admission_no).strip(), "payments": payments}
            )
    return rows
