"""Spreadsheet imports, the audit trail, and link handling."""
import io
from datetime import date

import openpyxl
import pytest

from conftest import login


def _sheet_bytes(wb):
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def test_payment_sheet_takes_several_installments(client, app, data):
    from app.models import Student

    login(client, "office1", "Office@123")
    sheet = client.get("/admin/payments/bulk-upload/template?installments=3")
    wb = openpyxl.load_workbook(io.BytesIO(sheet.data))
    ws = wb["Payments"]
    headers = [c.value for c in ws[1]]
    assert "installment_1" in headers and "installment_3" in headers

    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "S1")
    ws.cell(row=row, column=headers.index("installment_1") + 1, value=1000)
    ws.cell(row=row, column=headers.index("date_1") + 1, value="2026-06-30")
    ws.cell(row=row, column=headers.index("installment_2") + 1, value=500)
    ws.cell(row=row, column=headers.index("date_2") + 1, value="2026-07-31")

    client.post(
        "/admin/payments/bulk-upload",
        data={"file": (_sheet_bytes(wb), "p.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    with app.app_context():
        student = Student.query.get(data["student_a"])
        assert student.total_paid == 1500
        assert len(student.payments) == 2


def test_uploading_the_same_sheet_twice_does_not_double_payments(client, app, data):
    from app.models import Student

    login(client, "office1", "Office@123")
    sheet = client.get("/admin/payments/bulk-upload/template?installments=2")
    wb = openpyxl.load_workbook(io.BytesIO(sheet.data))
    ws = wb["Payments"]
    headers = [c.value for c in ws[1]]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "S1")
    ws.cell(row=row, column=headers.index("installment_1") + 1, value=750)
    ws.cell(row=row, column=headers.index("date_1") + 1, value="2026-06-30")

    for _ in range(2):
        client.post(
            "/admin/payments/bulk-upload",
            data={"file": (_sheet_bytes(wb), "p.xlsx")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )

    with app.app_context():
        assert Student.query.get(data["student_a"]).total_paid == 750


def test_payment_sheet_refuses_more_than_the_balance(client, app, data):
    from app.models import Student

    login(client, "office1", "Office@123")
    sheet = client.get("/admin/payments/bulk-upload/template?installments=2")
    wb = openpyxl.load_workbook(io.BytesIO(sheet.data))
    ws = wb["Payments"]
    headers = [c.value for c in ws[1]]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "S1")
    ws.cell(row=row, column=headers.index("installment_1") + 1, value=999999)

    client.post(
        "/admin/payments/bulk-upload",
        data={"file": (_sheet_bytes(wb), "p.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    with app.app_context():
        assert Student.query.get(data["student_a"]).total_paid == 0


def test_marks_sheet_rejects_out_of_range_values(client, app, data):
    from app.models import ExamMark

    login(client, "admin", "admin123")
    sheet = client.get(f"/admin/exams/{data['exam_id']}/marks/template")
    wb = openpyxl.load_workbook(io.BytesIO(sheet.data))
    ws = wb["Marks"]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "S1")
    ws.cell(row=row, column=4, value=500)

    client.post(
        f"/admin/exams/{data['exam_id']}/marks/upload",
        data={"file": (_sheet_bytes(wb), "m.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    with app.app_context():
        mark = ExamMark.query.filter_by(
            exam_subject_id=data["exam_subject_id"], student_id=data["student_a"]
        ).first()
        assert mark.marks_obtained == 90  # unchanged


def test_a_blank_marks_cell_leaves_the_mark_alone(client, app, data):
    from app.models import ExamMark

    login(client, "admin", "admin123")
    sheet = client.get(f"/admin/exams/{data['exam_id']}/marks/template")
    wb = openpyxl.load_workbook(io.BytesIO(sheet.data))
    ws = wb["Marks"]
    row = next(r for r in range(2, ws.max_row + 1) if ws.cell(row=r, column=1).value == "S1")
    ws.cell(row=row, column=4, value=None)

    client.post(
        f"/admin/exams/{data['exam_id']}/marks/upload",
        data={"file": (_sheet_bytes(wb), "m.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    with app.app_context():
        mark = ExamMark.query.filter_by(
            exam_subject_id=data["exam_subject_id"], student_id=data["student_a"]
        ).first()
        assert mark.marks_obtained == 90


def test_a_rubbish_file_is_rejected_politely(client, data):
    login(client, "office1", "Office@123")
    response = client.post(
        "/admin/payments/bulk-upload",
        data={"file": (io.BytesIO(b"not a spreadsheet"), "x.xlsx")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert b"Not a valid Excel" in response.data


def test_edits_are_written_to_the_audit_log(client, app, data):
    from app.models import AuditLog

    login(client, "office1", "Office@123")
    client.post(
        f"/admin/students/{data['student_a']}/payments/add",
        data={"amount": 100, "payment_date": date.today().isoformat(), "mode": "Cash"},
        follow_redirects=True,
    )
    with app.app_context():
        entry = AuditLog.query.filter_by(entity_type="payment").first()
        assert entry is not None
        assert entry.actor_name == "Office One"


def test_deleting_records_the_whole_row(client, app, data):
    from app.extensions import db
    from app.models import AuditLog, FeePayment

    login(client, "admin", "admin123")
    client.post(
        f"/admin/students/{data['student_a']}/payments/add",
        data={"amount": 100, "payment_date": date.today().isoformat(), "mode": "Cash"},
        follow_redirects=True,
    )
    with app.app_context():
        payment_id = FeePayment.query.first().id

    client.post(f"/admin/payments/{payment_id}/delete", follow_redirects=True)

    with app.app_context():
        entry = AuditLog.query.filter_by(action="deleted").first()
        assert entry is not None
        assert "amount" in (entry.summary or "")


def test_passwords_are_never_written_to_the_audit_log(client, app, data):
    from app.models import AuditLog

    login(client, "admin", "admin123")
    client.post(
        "/admin/office-staff",
        data={"name": "New Person", "username": "newoffice", "password": "Secret@999",
              "branch_id": data["high_branch"]},
        follow_redirects=True,
    )
    with app.app_context():
        for entry in AuditLog.query.all():
            assert "Secret@999" not in (entry.summary or "")


@pytest.mark.parametrize(
    "url,expected",
    [
        ("https://youtu.be/AZWyU5pkjL4", "youtube-nocookie.com/embed/AZWyU5pkjL4"),
        ("https://www.youtube.com/watch?v=AZWyU5pkjL4&t=9s", "embed/AZWyU5pkjL4"),
        ("https://vimeo.com/123456789", "player.vimeo.com/video/123456789"),
        ("https://youtube.com.evil.example/watch?v=AZWyU5pkjL4", None),
        ("https://example.com/whatever", None),
        ("javascript:alert(1)", None),
    ],
)
def test_only_known_video_hosts_are_embedded(url, expected):
    from app.utils.video import embed_url_for

    result = embed_url_for(url)
    if expected is None:
        assert result is None
    else:
        assert expected in result
